from fastapi import APIRouter, Depends, HTTPException, BackgroundTasks
from sqlalchemy.orm import Session
from typing import List, Optional, Dict, Any
from pydantic import BaseModel
from datetime import datetime, timezone
from io import BytesIO
import logging
import json
import os
import uuid
from pathlib import Path

from database.database import get_db
from auth.dependencies import get_current_organization, get_current_user
from auth.organization import ActiveOrganization
from auth.models import User
from database.models import RecordingSession

router = APIRouter(prefix="/api/export", tags=["Export"])
logger = logging.getLogger(__name__)

# Export models
class ExportOptions(BaseModel):
    includeTimestamps: bool = True
    includeSpeakers: bool = True
    includeInsights: bool = False
    mergeFiles: bool = False
    emailTo: Optional[str] = None
    scheduleTime: Optional[str] = None

class BatchExportRequest(BaseModel):
    sessionIds: List[str]
    format: str  # txt, pdf, docx, csv, json, srt
    options: ExportOptions

class ExportJob(BaseModel):
    id: str
    sessionIds: List[str]
    format: str
    status: str  # pending, processing, completed, failed
    progress: int
    createdAt: datetime
    completedAt: Optional[datetime] = None
    fileSize: Optional[int] = None
    downloadUrl: Optional[str] = None
    error: Optional[str] = None
    options: Dict[str, Any]

class ExportTemplate(BaseModel):
    id: str
    name: str
    format: str
    options: Dict[str, Any]
    description: str


def _get_session_for_export(
    db: Session,
    organization_id: int,
    session_identifier: str,
) -> Optional[RecordingSession]:
    session = db.query(RecordingSession).filter(
        RecordingSession.organization_id == organization_id,
        RecordingSession.session_id == session_identifier,
    ).first()
    if session:
        return session

    try:
        int_id = int(session_identifier)
    except (TypeError, ValueError):
        return None

    return db.query(RecordingSession).filter(
        RecordingSession.organization_id == organization_id,
        RecordingSession.id == int_id,
    ).first()

# In-memory storage for export jobs (should use database in production)
export_jobs: Dict[str, ExportJob] = {}
export_job_scope: Dict[str, Dict[str, int]] = {}
export_templates: List[ExportTemplate] = [
    ExportTemplate(
        id="1",
        name="Meeting Summary",
        format="pdf",
        options={"includeTimestamps": False, "includeSpeakers": True, "includeInsights": True},
        description="Professional meeting summary with key insights"
    ),
    ExportTemplate(
        id="2",
        name="Full Transcript",
        format="docx",
        options={"includeTimestamps": True, "includeSpeakers": True, "includeInsights": False},
        description="Complete transcript with timestamps and speakers"
    ),
    ExportTemplate(
        id="3",
        name="Subtitles",
        format="srt",
        options={"includeTimestamps": True, "includeSpeakers": False, "includeInsights": False},
        description="SRT subtitle format for video editing"
    )
]
export_template_scope: Dict[str, Dict[str, int]] = {}
BUILTIN_TEMPLATE_IDS = {template.id for template in export_templates}

@router.post("/batch", response_model=ExportJob)
async def create_batch_export(
    request: BatchExportRequest,
    background_tasks: BackgroundTasks,
    current_user: User = Depends(get_current_user),
    active_org: ActiveOrganization = Depends(get_current_organization),
    db: Session = Depends(get_db)
):
    """Create a batch export job for multiple sessions"""
    try:
        for session_id in request.sessionIds:
            if not _get_session_for_export(db, active_org.organization.id, session_id):
                raise HTTPException(
                    status_code=404,
                    detail="One or more sessions were not found in the active organization",
                )

        # Create export job
        job_id = str(uuid.uuid4())
        job = ExportJob(
            id=job_id,
            sessionIds=request.sessionIds,
            format=request.format,
            status="pending",
            progress=0,
            createdAt=datetime.now(timezone.utc),
            options=request.options.dict()
        )
        
        # Store job
        export_jobs[job_id] = job
        export_job_scope[job_id] = {
            "user_id": current_user.id,
            "organization_id": active_org.organization.id,
        }
        
        # Start background export task
        background_tasks.add_task(
            process_export_job,
            job_id,
            request.sessionIds,
            request.format,
            request.options,
            active_org.organization.id,
            db
        )
        
        logger.info(f"Created batch export job {job_id} for {len(request.sessionIds)} sessions")
        return job
        
    except Exception as e:
        logger.error(f"Failed to create batch export: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/jobs", response_model=List[ExportJob])
async def get_export_jobs(
    current_user: User = Depends(get_current_user),
    active_org: ActiveOrganization = Depends(get_current_organization),
):
    """Get all export jobs for the current user"""
    return [
        job
        for job_id, job in export_jobs.items()
        if export_job_scope.get(job_id, {}).get("user_id") == current_user.id
        and export_job_scope.get(job_id, {}).get("organization_id") == active_org.organization.id
    ]

@router.get("/jobs/{job_id}", response_model=ExportJob)
async def get_export_job(
    job_id: str,
    current_user: User = Depends(get_current_user),
    active_org: ActiveOrganization = Depends(get_current_organization),
):
    """Get status of a specific export job"""
    job = export_jobs.get(job_id)
    scope = export_job_scope.get(job_id, {})
    if not job or scope.get("user_id") != current_user.id or scope.get("organization_id") != active_org.organization.id:
        raise HTTPException(status_code=404, detail="Export job not found")
    return job

@router.delete("/jobs/{job_id}")
async def cancel_export_job(
    job_id: str,
    current_user: User = Depends(get_current_user),
    active_org: ActiveOrganization = Depends(get_current_organization),
):
    """Cancel an export job"""
    job = export_jobs.get(job_id)
    scope = export_job_scope.get(job_id, {})
    if not job or scope.get("user_id") != current_user.id or scope.get("organization_id") != active_org.organization.id:
        raise HTTPException(status_code=404, detail="Export job not found")
    
    if job.status in ["completed", "failed"]:
        raise HTTPException(status_code=400, detail="Cannot cancel completed or failed job")
    
    job.status = "cancelled"
    job.error = "Job cancelled by user"
    
    return {"message": "Export job cancelled"}

@router.get("/templates", response_model=List[ExportTemplate])
async def get_export_templates(
    current_user: User = Depends(get_current_user),
    active_org: ActiveOrganization = Depends(get_current_organization),
):
    """Get available export templates"""
    return [
        template
        for template in export_templates
        if template.id in BUILTIN_TEMPLATE_IDS
        or export_template_scope.get(template.id, {}).get("organization_id") == active_org.organization.id
    ]

@router.post("/templates", response_model=ExportTemplate)
async def create_export_template(
    template: ExportTemplate,
    current_user: User = Depends(get_current_user),
    active_org: ActiveOrganization = Depends(get_current_organization),
):
    """Create a new export template"""
    template.id = str(uuid.uuid4())
    export_templates.append(template)
    export_template_scope[template.id] = {
        "user_id": current_user.id,
        "organization_id": active_org.organization.id,
    }
    return template

@router.delete("/templates/{template_id}")
async def delete_export_template(
    template_id: str,
    current_user: User = Depends(get_current_user),
    active_org: ActiveOrganization = Depends(get_current_organization),
):
    """Delete an export template"""
    global export_templates
    scope = export_template_scope.get(template_id, {})
    if template_id not in BUILTIN_TEMPLATE_IDS and (
        scope.get("organization_id") != active_org.organization.id
        or scope.get("user_id") != current_user.id
    ):
        raise HTTPException(status_code=404, detail="Template not found")
    export_templates = [t for t in export_templates if t.id != template_id]
    export_template_scope.pop(template_id, None)
    return {"message": "Template deleted"}

async def process_export_job(
    job_id: str,
    session_ids: List[str],
    format: str,
    options: ExportOptions,
    organization_id: int,
    db: Session
):
    """Background task to process export job"""
    job = export_jobs.get(job_id)
    if not job:
        return
    
    try:
        job.status = "processing"
        total_sessions = len(session_ids)
        
        # Create export directory
        export_dir = Path("/tmp/exports") / job_id
        export_dir.mkdir(parents=True, exist_ok=True)
        
        exported_files = []
        
        for i, session_id in enumerate(session_ids):
            # Update progress
            job.progress = int((i / total_sessions) * 100)
            
            # Get session from database
            session = _get_session_for_export(db, organization_id, session_id)
            
            if not session:
                logger.warning(f"Session {session_id} not found")
                continue
            
            # Export based on format
            if format == "txt":
                content = export_to_txt(session, options)
                filename = f"{session.name or session_id}.txt"
            elif format == "json":
                content = export_to_json(session, options)
                filename = f"{session.name or session_id}.json"
            elif format == "srt":
                content = export_to_srt(session, options)
                filename = f"{session.name or session_id}.srt"
            elif format == "pdf":
                pdf_bytes = export_to_pdf(session, options)
                filename = f"{session.name or session_id}.pdf"
                file_path = export_dir / filename
                with open(file_path, 'wb') as f:
                    f.write(pdf_bytes)
                exported_files.append(file_path)
                continue
            elif format == "docx":
                docx_bytes = export_to_docx(session, options)
                filename = f"{session.name or session_id}.docx"
                file_path = export_dir / filename
                with open(file_path, 'wb') as f:
                    f.write(docx_bytes)
                exported_files.append(file_path)
                continue
            else:
                content = f"# {session.name or session_id}\n\nExport format {format} not supported."
                filename = f"{session.name or session_id}.txt"

            # Save text-based file
            file_path = export_dir / filename
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(content)
            
            exported_files.append(file_path)
        
        # Create archive if multiple files or merge requested
        if len(exported_files) > 1 or options.mergeFiles:
            import zipfile
            archive_path = export_dir / f"export_{job_id}.zip"
            with zipfile.ZipFile(archive_path, 'w') as zf:
                for file_path in exported_files:
                    zf.write(file_path, file_path.name)
            
            job.downloadUrl = f"/api/export/download/{job_id}/export_{job_id}.zip"
            job.fileSize = archive_path.stat().st_size
        elif exported_files:
            # Single file
            job.downloadUrl = f"/api/export/download/{job_id}/{exported_files[0].name}"
            job.fileSize = exported_files[0].stat().st_size
        
        job.status = "completed"
        job.progress = 100
        job.completedAt = datetime.now(timezone.utc)
        
        # Send email if requested
        if options.emailTo:
            # TODO: Implement email sending
            logger.info(f"Would send export to {options.emailTo}")
        
    except Exception as e:
        logger.error(f"Export job {job_id} failed: {e}")
        job.status = "failed"
        job.error = str(e)

def _get_summary_data(session: RecordingSession) -> Dict[str, Any]:
    """Extract summary data from session, handling both legacy and new formats."""
    summary_data = {
        "executive": "",
        "bullets": [],
        "actions": [],
        "decisions": [],
        "title": session.title or session.name or "Untitled Meeting",
    }

    # Try final_summary first (preferred)
    final = session.final_summary
    if final:
        if isinstance(final, str):
            try:
                final = json.loads(final)
            except (json.JSONDecodeError, TypeError):
                final = {}
        if isinstance(final, dict):
            summary_data["executive"] = final.get("executive", "")
            summary_data["bullets"] = final.get("bullets", [])
            summary_data["actions"] = final.get("actions", [])
            summary_data["decisions"] = final.get("decisions", [])
            if final.get("title"):
                summary_data["title"] = final["title"]
            return summary_data

    # Fallback to legacy summary field
    legacy = session.summary
    if legacy:
        if isinstance(legacy, str):
            try:
                legacy = json.loads(legacy)
            except (json.JSONDecodeError, TypeError):
                legacy = {}
        if isinstance(legacy, dict):
            # Could be nested under 'analysis' key
            source = legacy.get("analysis", legacy)
            summary_data["executive"] = source.get("executive", "")
            summary_data["bullets"] = source.get("bullets", [])
            summary_data["actions"] = source.get("actions", [])
            summary_data["decisions"] = source.get("decisions", [])
            if source.get("title"):
                summary_data["title"] = source["title"]

    return summary_data


def _get_transcript_segments(session: RecordingSession) -> list:
    """Extract transcript segments from session data."""
    # Try diarized transcript first
    if session.transcript_diarized and isinstance(session.transcript_diarized, dict):
        segments = session.transcript_diarized.get("segments", [])
        if segments:
            return segments

    # Try transcript JSON field
    if session.transcript:
        try:
            data = json.loads(session.transcript) if isinstance(session.transcript, str) else session.transcript
            if isinstance(data, dict) and "segments" in data:
                return data["segments"]
        except (json.JSONDecodeError, TypeError):
            pass

    return []


def _get_plain_transcript(session: RecordingSession) -> str:
    """Get plain text transcript from session."""
    if session.transcript_simple:
        return session.transcript_simple
    if session.transcript:
        try:
            data = json.loads(session.transcript) if isinstance(session.transcript, str) else session.transcript
            if isinstance(data, dict) and "text" in data:
                return data["text"]
            if isinstance(data, str):
                return data
        except (json.JSONDecodeError, TypeError):
            return str(session.transcript)
    return ""


def export_to_pdf(session: RecordingSession, options: ExportOptions) -> bytes:
    """Export session to a professional PDF document using reportlab."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        HRFlowable, PageBreak
    )
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT

    buffer = BytesIO()
    page_width, page_height = letter
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        topMargin=1.0 * inch,
        bottomMargin=0.8 * inch,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
    )

    # Colors
    purple = HexColor("#7C3AED")
    dark_gray = HexColor("#374151")
    medium_gray = HexColor("#6B7280")
    light_gray = HexColor("#F3F4F6")
    yellow_bg = HexColor("#FEF3C7")
    green_text = HexColor("#065F46")

    # Custom styles
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        "MeetingTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        textColor=dark_gray,
        spaceAfter=4,
    ))
    styles.add(ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=14,
        textColor=purple,
        spaceBefore=16,
        spaceAfter=8,
    ))
    styles.add(ParagraphStyle(
        "BodyText2",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        textColor=dark_gray,
        leading=14,
        spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        "BulletText",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        textColor=dark_gray,
        leading=14,
        leftIndent=16,
        spaceAfter=4,
        bulletIndent=0,
    ))
    styles.add(ParagraphStyle(
        "MetaText",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        textColor=medium_gray,
        spaceAfter=2,
    ))
    styles.add(ParagraphStyle(
        "TranscriptSpeaker",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=9,
        textColor=purple,
        spaceAfter=1,
    ))
    styles.add(ParagraphStyle(
        "TranscriptText",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        textColor=dark_gray,
        leading=12,
        spaceAfter=8,
        leftIndent=8,
    ))

    story = []
    summary_data = _get_summary_data(session)

    # --- Header ---
    story.append(Paragraph("Meeting-Ops", ParagraphStyle(
        "BrandName", parent=styles["Normal"],
        fontName="Helvetica-Bold", fontSize=12, textColor=purple,
    )))
    story.append(HRFlowable(
        width="100%", thickness=2, color=purple, spaceAfter=12,
    ))
    story.append(Paragraph(
        summary_data["title"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"),
        styles["MeetingTitle"]
    ))

    # Meeting metadata
    date_str = session.created_at.strftime("%B %d, %Y at %I:%M %p") if session.created_at else "N/A"
    duration_secs = session.duration or 0
    duration_str = f"{int(duration_secs // 60)}m {int(duration_secs % 60)}s" if duration_secs else "N/A"
    story.append(Paragraph(f"Date: {date_str}", styles["MetaText"]))
    story.append(Paragraph(f"Duration: {duration_str}", styles["MetaText"]))
    story.append(Spacer(1, 12))

    has_summary = bool(
        summary_data["executive"] or summary_data["bullets"]
        or summary_data["actions"] or summary_data["decisions"]
    )

    # --- Executive Summary ---
    if summary_data["executive"]:
        story.append(Paragraph("Executive Summary", styles["SectionHeading"]))
        safe_text = summary_data["executive"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(safe_text, styles["BodyText2"]))
        story.append(Spacer(1, 8))

    # --- Key Discussion Points ---
    if summary_data["bullets"]:
        story.append(Paragraph("Key Discussion Points", styles["SectionHeading"]))
        for bullet in summary_data["bullets"]:
            safe = str(bullet).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(f"\u2022  {safe}", styles["BulletText"]))
        story.append(Spacer(1, 8))

    # --- Action Items ---
    if summary_data["actions"]:
        story.append(Paragraph("Action Items", styles["SectionHeading"]))
        for action in summary_data["actions"]:
            if isinstance(action, dict):
                action_text = action.get("action", str(action))
                owner = action.get("owner", "")
                priority = action.get("priority", "")
                safe_action = action_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                line = f"\u25a1  {safe_action}"
                if owner:
                    line += f"  &mdash; <i>{owner.replace('&', '&amp;')}</i>"
                if priority:
                    line += f"  [{priority}]"
                story.append(Paragraph(line, styles["BulletText"]))
            else:
                safe = str(action).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(f"\u25a1  {safe}", styles["BulletText"]))
        story.append(Spacer(1, 8))

    # --- Important Decisions ---
    if summary_data["decisions"]:
        story.append(Paragraph("Important Decisions", styles["SectionHeading"]))
        for decision in summary_data["decisions"]:
            safe = str(decision).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(f"\u2713  {safe}", styles["BulletText"]))
        story.append(Spacer(1, 8))

    # --- Transcript ---
    segments = _get_transcript_segments(session)
    plain_transcript = _get_plain_transcript(session)

    if segments or plain_transcript:
        if has_summary:
            story.append(PageBreak())
        story.append(Paragraph("Full Transcript", styles["SectionHeading"]))

        if segments:
            for seg in segments:
                speaker = seg.get("speaker", "Speaker")
                text = seg.get("text", "")
                start = seg.get("start", 0)

                label_parts = []
                if options.includeSpeakers:
                    safe_speaker = speaker.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    label_parts.append(safe_speaker)
                if options.includeTimestamps:
                    label_parts.append(f"[{format_timestamp(start)}]")

                if label_parts:
                    story.append(Paragraph(" ".join(label_parts), styles["TranscriptSpeaker"]))

                safe_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe_text, styles["TranscriptText"]))
        elif plain_transcript:
            # Wrap long plain text into paragraphs
            for para in plain_transcript.split("\n"):
                if para.strip():
                    safe = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    story.append(Paragraph(safe, styles["TranscriptText"]))
    elif not has_summary:
        story.append(Paragraph("No transcript or summary data available for this session.", styles["BodyText2"]))

    # --- Footer callback ---
    def add_footer(canvas, doc):
        canvas.saveState()
        # Purple accent line
        canvas.setStrokeColor(purple)
        canvas.setLineWidth(0.5)
        canvas.line(0.75 * inch, 0.6 * inch, page_width - 0.75 * inch, 0.6 * inch)
        # Footer text
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(medium_gray)
        canvas.drawString(0.75 * inch, 0.45 * inch, "Generated by Meeting-Ops AI Recording Suite")
        canvas.drawRightString(page_width - 0.75 * inch, 0.45 * inch, f"Page {doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=add_footer, onLaterPages=add_footer)
    return buffer.getvalue()


def export_to_docx(session: RecordingSession, options: ExportOptions) -> bytes:
    """Export session to a professional DOCX document using python-docx."""
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    # --- Style configuration ---
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(10)
    font.color.rgb = RGBColor(0x37, 0x41, 0x51)

    purple_rgb = RGBColor(0x7C, 0x3A, 0xED)
    dark_gray_rgb = RGBColor(0x37, 0x41, 0x51)
    medium_gray_rgb = RGBColor(0x6B, 0x72, 0x80)

    summary_data = _get_summary_data(session)

    # --- Brand header ---
    brand_para = doc.add_paragraph()
    brand_run = brand_para.add_run("Meeting-Ops")
    brand_run.font.size = Pt(12)
    brand_run.font.bold = True
    brand_run.font.color.rgb = purple_rgb

    # Purple divider line (thin horizontal rule via a border paragraph)
    divider = doc.add_paragraph()
    divider_fmt = divider.paragraph_format
    divider_fmt.space_after = Pt(6)
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "7C3AED")
    pBdr.append(bottom)
    divider.paragraph_format.element.get_or_add_pPr().append(pBdr)

    # --- Title ---
    title_para = doc.add_heading(summary_data["title"], level=1)
    for run in title_para.runs:
        run.font.color.rgb = dark_gray_rgb

    # --- Metadata ---
    date_str = session.created_at.strftime("%B %d, %Y at %I:%M %p") if session.created_at else "N/A"
    duration_secs = session.duration or 0
    duration_str = f"{int(duration_secs // 60)}m {int(duration_secs % 60)}s" if duration_secs else "N/A"

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(2)
    date_run = meta.add_run(f"Date: {date_str}")
    date_run.font.size = Pt(9)
    date_run.font.color.rgb = medium_gray_rgb
    meta.add_run("\n")
    dur_run = meta.add_run(f"Duration: {duration_str}")
    dur_run.font.size = Pt(9)
    dur_run.font.color.rgb = medium_gray_rgb

    has_summary = bool(
        summary_data["executive"] or summary_data["bullets"]
        or summary_data["actions"] or summary_data["decisions"]
    )

    # --- Executive Summary ---
    if summary_data["executive"]:
        doc.add_heading("Executive Summary", level=2)
        doc.add_paragraph(summary_data["executive"])

    # --- Key Discussion Points ---
    if summary_data["bullets"]:
        doc.add_heading("Key Discussion Points", level=2)
        for bullet in summary_data["bullets"]:
            doc.add_paragraph(str(bullet), style="List Bullet")

    # --- Action Items (Table) ---
    if summary_data["actions"]:
        doc.add_heading("Action Items", level=2)

        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["Task", "Owner", "Priority", "Status"]
        header_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            header_cells[i].text = header_text
            for paragraph in header_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.size = Pt(9)

        for action in summary_data["actions"]:
            if isinstance(action, dict):
                task = action.get("action", str(action))
                owner = action.get("owner", "")
                priority = action.get("priority", "")
                status = action.get("status", "Pending")
            else:
                task = str(action)
                owner = ""
                priority = ""
                status = "Pending"

            row_cells = table.add_row().cells
            row_cells[0].text = task
            row_cells[1].text = owner
            row_cells[2].text = priority
            row_cells[3].text = status

            for cell in row_cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)

    # --- Important Decisions ---
    if summary_data["decisions"]:
        doc.add_heading("Important Decisions", level=2)
        for decision in summary_data["decisions"]:
            doc.add_paragraph(str(decision), style="List Bullet")

    # --- Full Transcript ---
    segments = _get_transcript_segments(session)
    plain_transcript = _get_plain_transcript(session)

    if segments or plain_transcript:
        if has_summary:
            doc.add_page_break()
        doc.add_heading("Full Transcript", level=2)

        if segments:
            for seg in segments:
                speaker = seg.get("speaker", "Speaker")
                text = seg.get("text", "")
                start = seg.get("start", 0)

                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)

                label_parts = []
                if options.includeSpeakers:
                    label_parts.append(speaker)
                if options.includeTimestamps:
                    label_parts.append(f"[{format_timestamp(start)}]")

                if label_parts:
                    label_run = p.add_run(" ".join(label_parts) + ": ")
                    label_run.font.bold = True
                    label_run.font.size = Pt(9)
                    label_run.font.color.rgb = purple_rgb

                text_run = p.add_run(text)
                text_run.font.size = Pt(9)
        elif plain_transcript:
            for para_text in plain_transcript.split("\n"):
                if para_text.strip():
                    p = doc.add_paragraph(para_text)
                    p.paragraph_format.space_after = Pt(4)
                    for run in p.runs:
                        run.font.size = Pt(9)
    elif not has_summary:
        doc.add_paragraph("No transcript or summary data available for this session.")

    # --- Footer ---
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run("Generated by Meeting-Ops AI Recording Suite")
    footer_run.font.size = Pt(7)
    footer_run.font.color.rgb = medium_gray_rgb

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_to_txt(session: RecordingSession, options: ExportOptions) -> str:
    """Export session to plain text format"""
    lines = []
    lines.append(f"Meeting: {session.name or 'Untitled'}")
    lines.append(f"Date: {session.created_at}")
    lines.append(f"Duration: {session.duration or 0} seconds")
    lines.append("-" * 50)
    lines.append("")
    
    # Add transcript if available
    if hasattr(session, 'transcript') and session.transcript:
        transcript = json.loads(session.transcript) if isinstance(session.transcript, str) else session.transcript
        
        for segment in transcript.get('segments', []):
            line = ""
            
            if options.includeTimestamps:
                start_time = segment.get('start', 0)
                line += f"[{format_timestamp(start_time)}] "
            
            if options.includeSpeakers and 'speaker' in segment:
                line += f"{segment['speaker']}: "
            
            line += segment.get('text', '')
            lines.append(line)
    else:
        lines.append("No transcript available for this session.")
    
    if options.includeInsights:
        lines.append("")
        lines.append("-" * 50)
        lines.append("AI INSIGHTS")
        lines.append("-" * 50)
        lines.append("• Key topics discussed")
        lines.append("• Action items identified")
        lines.append("• Decisions made")
    
    return "\n".join(lines)

def export_to_json(session: RecordingSession, options: ExportOptions) -> str:
    """Export session to JSON format"""
    data = {
        "session_id": session.id,
        "name": session.name,
        "created_at": session.created_at.isoformat() if session.created_at else None,
        "duration": session.duration,
        "transcript": []
    }
    
    if hasattr(session, 'transcript') and session.transcript:
        transcript = json.loads(session.transcript) if isinstance(session.transcript, str) else session.transcript
        
        for segment in transcript.get('segments', []):
            segment_data = {"text": segment.get('text', '')}
            
            if options.includeTimestamps:
                segment_data["start"] = segment.get('start', 0)
                segment_data["end"] = segment.get('end', 0)
            
            if options.includeSpeakers and 'speaker' in segment:
                segment_data["speaker"] = segment['speaker']
            
            data["transcript"].append(segment_data)
    
    if options.includeInsights:
        data["insights"] = {
            "topics": [],
            "action_items": [],
            "decisions": []
        }
    
    return json.dumps(data, indent=2)

def export_to_srt(session: RecordingSession, options: ExportOptions) -> str:
    """Export session to SRT subtitle format"""
    lines = []
    counter = 1
    
    if hasattr(session, 'transcript') and session.transcript:
        transcript = json.loads(session.transcript) if isinstance(session.transcript, str) else session.transcript
        
        for segment in transcript.get('segments', []):
            lines.append(str(counter))
            
            start_time = format_srt_timestamp(segment.get('start', 0))
            end_time = format_srt_timestamp(segment.get('end', segment.get('start', 0) + 5))
            lines.append(f"{start_time} --> {end_time}")
            
            text = segment.get('text', '')
            if options.includeSpeakers and 'speaker' in segment:
                text = f"{segment['speaker']}: {text}"
            
            lines.append(text)
            lines.append("")  # Empty line between subtitles
            counter += 1
    
    return "\n".join(lines)

def format_timestamp(seconds: float) -> str:
    """Format seconds to HH:MM:SS"""
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"

def format_srt_timestamp(seconds: float) -> str:
    """Format seconds to SRT timestamp format (HH:MM:SS,mmm)"""
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = seconds % 60
    return f"{hours:02d}:{minutes:02d}:{secs:06.3f}".replace('.', ',')

@router.get("/download/{job_id}/{filename}")
async def download_export(
    job_id: str,
    filename: str,
    current_user: User = Depends(get_current_user),
    active_org: ActiveOrganization = Depends(get_current_organization),
):
    """Download exported file"""
    from fastapi.responses import FileResponse
    
    scope = export_job_scope.get(job_id, {})
    if scope.get("user_id") != current_user.id or scope.get("organization_id") != active_org.organization.id:
        raise HTTPException(status_code=404, detail="Export job not found")

    file_path = Path("/tmp/exports") / job_id / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    
    return FileResponse(
        path=str(file_path),
        filename=filename,
        media_type='application/octet-stream'
    )
